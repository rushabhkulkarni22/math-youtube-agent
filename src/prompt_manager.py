from pathlib import Path

from openpyxl import load_workbook

from src.database import StateDatabase


HEADERS = [
    "Prompt ID", "Status", "Video File", "Thumbnail File", "YouTube Video ID",
    "YouTube URL", "Upload Date", "Retry Count", "Error Message",
]


class PromptManager:
    def __init__(self, workbook_path: Path, database: StateDatabase):
        self.workbook_path = workbook_path
        self.database = database

    def import_prompts(self) -> int:
        workbook = load_workbook(self.workbook_path)
        sheet = workbook.active
        for column, header in enumerate(HEADERS, start=2):
            sheet.cell(row=1, column=column, value=header)
        count = 0
        for excel_row in range(2, sheet.max_row + 1):
            prompt = sheet.cell(excel_row, 1).value
            if not prompt or not str(prompt).strip():
                continue
            prompt_id = excel_row - 1
            self.database.upsert_prompt(prompt_id, str(prompt).strip(), excel_row)
            sheet.cell(excel_row, 2, prompt_id)
            if not sheet.cell(excel_row, 3).value:
                sheet.cell(excel_row, 3, "PENDING")
            count += 1
        workbook.save(self.workbook_path)
        return count

    def update_excel(self, prompt_id: int, status: str, **values: object) -> None:
        prompt = self.database.get_prompt(prompt_id)
        if not prompt or not prompt["excel_row"]:
            return
        workbook = load_workbook(self.workbook_path)
        sheet = workbook.active
        row = int(prompt["excel_row"])
        mapping = {
            "status": 3, "video_file": 4, "thumbnail_file": 5,
            "youtube_video_id": 6, "youtube_url": 7, "upload_date": 8,
            "retry_count": 9, "error_message": 10,
        }
        sheet.cell(row, 3, status)
        for key, value in values.items():
            if key in mapping:
                sheet.cell(row, mapping[key], value)
        workbook.save(self.workbook_path)
